"""
Document output tools: turn an employee's finished work into a file.

Every file is written to WORKSPACE_DIR and, when called as an agent tool, also
saved as a session artifact so the Chief of Staff (and the user) receive the
actual file, not just a server-local path.
"""

import asyncio
import inspect
import logging
import re
from datetime import datetime
from pathlib import Path

from fpdf import FPDF
from google.adk.agents.llm_agent import ToolUnion
from google.adk.tools.tool_context import ToolContext
from google.genai import types as genai_types

logger = logging.getLogger(__name__)

WORKSPACE_DIR = Path(__file__).parent.parent / "workspace"

_MAX_TABLE_ROWS = 200


def _latin1(text: object) -> str:
    """fpdf's core fonts are Latin-1 only; replace anything else with '?'
    rather than failing the whole document."""
    return str(text).encode("latin-1", "replace").decode("latin-1")


def _safe_stem(name: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", name.strip()).strip("_").lower()
    return stem[:60] or "report"


class ReportPDF(FPDF):
    """A plain A4 report: title on every page, page number in the footer."""

    def __init__(self, title: str):
        super().__init__()
        self.report_title = _latin1(title)
        self.set_auto_page_break(auto=True, margin=15)

    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 10, self.report_title, align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(3)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(
            0, 10,
            f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')} | Page {self.page_no()}",
            align="C",
        )

    def add_heading(self, text: str):
        self.set_font("Helvetica", "B", 12)
        self.set_fill_color(240, 240, 240)
        self.cell(0, 9, _latin1(text), fill=True, new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def add_text(self, text: str):
        self.set_font("Helvetica", "", 10)
        self.multi_cell(0, 6, _latin1(text), new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def add_table(self, headers: list, rows: list):
        if not headers:
            return
        col_width = (self.w - 20) / len(headers)
        self.set_font("Helvetica", "B", 9)
        self.set_fill_color(210, 210, 210)
        for h in headers:
            self.cell(col_width, 7, _latin1(h)[:30], border=1, fill=True, align="C")
        self.ln()
        self.set_font("Helvetica", "", 8)
        for row in rows[:_MAX_TABLE_ROWS]:
            for i in range(len(headers)):
                cell = row[i] if i < len(row) else ""
                self.cell(col_width, 6, _latin1(cell)[:35], border=1)
            self.ln()
        self.ln(3)


async def _save_file(
    data: bytes, filename: str, mime: str, tool_context: ToolContext | None
) -> str:
    """Write `data` to WORKSPACE_DIR and, when called as a tool, save it as an
    artifact too. Returns the on-disk path."""
    WORKSPACE_DIR.mkdir(exist_ok=True)
    filepath = WORKSPACE_DIR / filename
    # Blocking disk write off the event loop: these tools are coroutines and
    # must not stall the employee's server (a stalled loop fails health checks).
    await asyncio.to_thread(filepath.write_bytes, data)
    if tool_context is not None:
        try:
            part = genai_types.Part(inline_data=genai_types.Blob(data=data, mime_type=mime))
            result = tool_context.save_artifact(filename, part)
            if inspect.isawaitable(result):
                await result
        except Exception as e:  # artifact service missing / backend error
            logger.warning("could not save artifact %s: %s", filename, e)
    return str(filepath)


async def write_pdf_report(
    title: str,
    sections: list[dict],
    file_name: str | None = None,
    tool_context: ToolContext | None = None,
) -> dict:
    """
    Write a PDF report and return it to the user as a downloadable file.

    Args:
        title: the report title, shown at the top of every page.
        sections: the report body, in order. Each section is a dict with any of:
            "heading" (str), "text" (str, paragraphs separated by blank lines),
            "bullets" ([str, ...]), "table" ({"headers": [str, ...],
            "rows": [[value, ...], ...]}). Tables show at most 200 rows.
        file_name: optional base name (no extension). Defaults to the title.

    Returns:
        dict: success, file_name, file_path, sections; on failure success=False
            and an error.
    """
    try:
        if not isinstance(sections, list) or not sections:
            return {"success": False, "error": "sections must be a non-empty list."}
        pdf = ReportPDF(title)
        pdf.add_page()
        for section in sections:
            if not isinstance(section, dict):
                continue
            if section.get("heading"):
                pdf.add_heading(section["heading"])
            if section.get("text"):
                pdf.add_text(section["text"])
            for item in section.get("bullets") or []:
                pdf.add_text(f"- {item}")
            table = section.get("table")
            if isinstance(table, dict):
                pdf.add_table(table.get("headers") or [], table.get("rows") or [])

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{_safe_stem(file_name or title)}_{stamp}.pdf"
        filepath = await _save_file(
            bytes(pdf.output()), filename, "application/pdf", tool_context
        )
        return {
            "success": True,
            "file_name": filename,
            "file_path": filepath,
            "sections": len(sections),
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def _stamped(base: str, ext: str) -> str:
    return f"{_safe_stem(base)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"


_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def write_word_document(
    title: str,
    sections: list[dict],
    file_name: str | None = None,
    tool_context: ToolContext | None = None,
) -> dict:
    """
    Write a Word document (.docx) - a proposal, letter, article or brief - and
    return it to the user as a downloadable file they can edit.

    Args:
        title: the document title.
        sections: the body, in order. Each section is a dict with any of:
            "heading" (str), "text" (str, paragraphs separated by blank lines),
            "bullets" ([str, ...]), "table" ({"headers": [...], "rows": [[...], ...]}).
        file_name: optional base name (no extension). Defaults to the title.

    Returns:
        dict: success, file_name, file_path, sections; on failure success=False
            and an error.
    """
    try:
        import io

        from docx import Document

        if not isinstance(sections, list) or not sections:
            return {"success": False, "error": "sections must be a non-empty list."}
        doc = Document()
        doc.add_heading(str(title), level=0)
        for section in sections:
            if not isinstance(section, dict):
                continue
            if section.get("heading"):
                doc.add_heading(str(section["heading"]), level=1)
            for para in str(section.get("text") or "").split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
            for item in section.get("bullets") or []:
                doc.add_paragraph(str(item), style="List Bullet")
            table = section.get("table")
            if isinstance(table, dict) and table.get("headers"):
                headers = [str(h) for h in table["headers"]]
                grid = doc.add_table(rows=1, cols=len(headers))
                grid.style = "Table Grid"
                for cell, h in zip(grid.rows[0].cells, headers):
                    cell.text = h
                for row in (table.get("rows") or [])[:_MAX_TABLE_ROWS]:
                    cells = grid.add_row().cells
                    for i in range(len(headers)):
                        cells[i].text = str(row[i]) if i < len(row) else ""
        buf = io.BytesIO()
        doc.save(buf)
        filename = _stamped(file_name or title, "docx")
        filepath = await _save_file(buf.getvalue(), filename, _DOCX_MIME, tool_context)
        return {"success": True, "file_name": filename, "file_path": filepath,
                "sections": len(sections)}
    except Exception as e:
        return {"success": False, "error": str(e)}


async def write_spreadsheet(
    sheets: list[dict],
    file_name: str,
    tool_context: ToolContext | None = None,
) -> dict:
    """
    Write an Excel workbook (.xlsx) - a lead list, a price table, a cleaned
    dataset - and return it to the user as a downloadable file.

    Args:
        sheets: one dict per sheet: {"name": str, "headers": [str, ...],
            "rows": [[value, ...], ...]}. Numbers stay numbers. At most 10,000
            rows per sheet.
        file_name: base name (no extension), e.g. "uae_integrators".

    Returns:
        dict: success, file_name, file_path, sheets (name and row count);
            on failure success=False and an error.
    """
    try:
        import io

        from openpyxl import Workbook
        from openpyxl.styles import Font

        if not isinstance(sheets, list) or not sheets:
            return {"success": False, "error": "sheets must be a non-empty list."}
        wb = Workbook()
        wb.remove(wb.active)
        summary = []
        for i, sheet in enumerate(sheets):
            if not isinstance(sheet, dict):
                continue
            name = re.sub(r"[\[\]:*?/\\]", "_", str(sheet.get("name") or f"Sheet{i + 1}"))[:31]
            ws = wb.create_sheet(name)
            headers = list(sheet.get("headers") or [])
            if headers:
                ws.append([str(h) for h in headers])
                for cell in ws[1]:
                    cell.font = Font(bold=True)
            rows = (sheet.get("rows") or [])[:10_000]
            for row in rows:
                ws.append([v if isinstance(v, (int, float)) or v is None else str(v)
                           for v in row])
            summary.append({"name": name, "rows": len(rows)})
        if not summary:
            return {"success": False, "error": "no valid sheets."}
        buf = io.BytesIO()
        wb.save(buf)
        filename = _stamped(file_name, "xlsx")
        filepath = await _save_file(buf.getvalue(), filename, _XLSX_MIME, tool_context)
        return {"success": True, "file_name": filename, "file_path": filepath, "sheets": summary}
    except Exception as e:
        return {"success": False, "error": str(e)}


_TEXT_TYPES = {"md": "text/markdown", "txt": "text/plain", "csv": "text/csv", "html": "text/html"}


async def write_text_file(
    file_name: str,
    content: str,
    file_type: str = "md",
    tool_context: ToolContext | None = None,
) -> dict:
    """
    Save text as a downloadable file: Markdown, plain text, CSV or HTML (for
    example a blog post, an email draft or a social media post).

    Args:
        file_name: base name (no extension).
        content: the full text of the file.
        file_type: "md" (default), "txt", "csv" or "html".

    Returns:
        dict: success, file_name, file_path, characters; on failure success=False
            and an error.
    """
    try:
        ext = (file_type or "md").lower().lstrip(".")
        if ext not in _TEXT_TYPES:
            return {"success": False,
                    "error": f"file_type must be one of: {', '.join(_TEXT_TYPES)}."}
        if not isinstance(content, str) or not content.strip():
            return {"success": False, "error": "content is empty."}
        filename = _stamped(file_name, ext)
        filepath = await _save_file(content.encode("utf-8"), filename, _TEXT_TYPES[ext],
                                    tool_context)
        return {"success": True, "file_name": filename, "file_path": filepath,
                "characters": len(content)}
    except Exception as e:
        return {"success": False, "error": str(e)}


def list_workspace_files() -> dict:
    """
    List the files employees have produced in the workspace, newest first.

    Returns:
        dict: success, total, and files - a list of {file_name, size_kb, created_at}.
    """
    try:
        files = []
        if WORKSPACE_DIR.exists():
            for f in WORKSPACE_DIR.iterdir():
                if f.is_file():
                    st = f.stat()
                    files.append({
                        "file_name": f.name,
                        "size_kb": round(st.st_size / 1024, 2),
                        "created_at": datetime.fromtimestamp(st.st_mtime).isoformat(),
                    })
        files.sort(key=lambda x: x["created_at"], reverse=True)
        return {"success": True, "files": files, "total": len(files)}
    except Exception as e:
        return {"success": False, "error": str(e)}


DOCUMENT_TOOLS: list[ToolUnion] = [
    write_pdf_report,
    write_word_document,
    write_spreadsheet,
    write_text_file,
    list_workspace_files,
]

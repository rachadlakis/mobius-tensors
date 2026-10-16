"""Tests for the document output tools in shared/documents.py.

PDFs are really written (to a tmp dir) so the fpdf layout code is exercised
end to end.
"""

import os
import time

import pytest

from shared import documents

SECTIONS = [
    {"heading": "Summary", "text": "Revenue grew 12%.\n\nCosts were flat."},
    {"heading": "By region", "table": {
        "headers": ["Region", "Revenue"],
        "rows": [["North", 120], ["South", 80]],
    }},
]


@pytest.fixture
def workspace(monkeypatch, tmp_path):
    monkeypatch.setattr(documents, "WORKSPACE_DIR", tmp_path)
    return tmp_path


def _assert_pdf_written(result, workspace):
    assert result["success"] is True
    path = workspace / result["file_name"]
    assert path.exists()
    assert path.read_bytes().startswith(b"%PDF")


class TestWritePdfReport:
    async def test_writes_pdf(self, workspace):
        result = await documents.write_pdf_report("Q3 Sales", SECTIONS)
        _assert_pdf_written(result, workspace)
        assert result["file_name"].startswith("q3_sales_")
        assert result["sections"] == 2

    async def test_custom_file_name_is_sanitised(self, workspace):
        result = await documents.write_pdf_report("T", SECTIONS, file_name="../Board pack!")
        _assert_pdf_written(result, workspace)
        assert result["file_name"].startswith("board_pack_")

    async def test_non_latin_text_does_not_fail(self, workspace):
        result = await documents.write_pdf_report(
            "تقرير", [{"heading": "ملخص", "text": "€ 100 — ok"}]
        )
        _assert_pdf_written(result, workspace)

    async def test_ragged_table_rows(self, workspace):
        result = await documents.write_pdf_report("T", [{"table": {
            "headers": ["a", "b", "c"], "rows": [[1], [1, 2, 3, 4]],
        }}])
        _assert_pdf_written(result, workspace)

    async def test_empty_sections_is_an_error(self, workspace):
        result = await documents.write_pdf_report("T", [])
        assert result["success"] is False


class TestArtifactSaving:
    async def test_saves_pdf_as_artifact_when_tool_context_present(self, workspace):
        saved = {}

        class FakeCtx:
            def save_artifact(self, filename, part):
                saved["filename"] = filename
                saved["bytes"] = part.inline_data.data
                saved["mime"] = part.inline_data.mime_type
                return 1  # sync, like a stub artifact service

        result = await documents.write_pdf_report("T", SECTIONS, tool_context=FakeCtx())
        assert result["success"] is True
        assert saved["filename"] == result["file_name"]
        assert saved["bytes"].startswith(b"%PDF")
        assert saved["mime"] == "application/pdf"

    async def test_artifact_save_failure_does_not_break_report(self, workspace):
        class BoomCtx:
            def save_artifact(self, filename, part):
                raise RuntimeError("no artifact service")

        result = await documents.write_pdf_report("T", SECTIONS, tool_context=BoomCtx())
        assert result["success"] is True  # PDF still written to disk

    async def test_async_save_artifact_is_awaited(self, workspace):
        seen = {}

        class AsyncCtx:
            async def save_artifact(self, filename, part):  # like the real ADK ToolContext
                seen["filename"] = filename
                return 1

        result = await documents.write_pdf_report("T", SECTIONS, tool_context=AsyncCtx())
        assert result["success"] is True
        assert seen["filename"] == result["file_name"]


class TestListWorkspaceFiles:
    def test_lists_files_newest_first(self, workspace):
        old = workspace / "old.pdf"
        new = workspace / "new.csv"
        old.write_bytes(b"%PDF-old")
        new.write_bytes(b"a,b")
        os.utime(old, (time.time() - 100, time.time() - 100))

        result = documents.list_workspace_files()
        assert result["success"] is True
        assert result["total"] == 2
        assert result["files"][0]["file_name"] == "new.csv"
        assert "size_kb" in result["files"][0]

    def test_missing_dir(self, monkeypatch, tmp_path):
        monkeypatch.setattr(documents, "WORKSPACE_DIR", tmp_path / "nope")
        assert documents.list_workspace_files() == {"success": True, "files": [], "total": 0}


class TestWordDocument:
    async def test_writes_docx_with_all_section_kinds(self, workspace):
        from docx import Document

        result = await documents.write_word_document("Proposal for Acme", [
            {"heading": "Offer", "text": "Para one.\n\nPara two.", "bullets": ["fast", "cheap"]},
            {"table": {"headers": ["Item", "Price"], "rows": [["Desk", 420], ["Chair"]]}},
        ])
        assert result["success"] is True
        assert result["file_name"].endswith(".docx")
        doc = Document(str(workspace / result["file_name"]))
        texts = [p.text for p in doc.paragraphs]
        assert "Proposal for Acme" in texts
        assert "Para two." in texts
        assert "cheap" in texts
        assert doc.tables[0].rows[2].cells[0].text == "Chair"

    async def test_empty_sections(self, workspace):
        assert (await documents.write_word_document("T", []))["success"] is False


class TestSpreadsheet:
    async def test_writes_sheets_keeping_numbers(self, workspace):
        from openpyxl import load_workbook

        result = await documents.write_spreadsheet([
            {"name": "Leads: UAE", "headers": ["company", "size"],
             "rows": [["Acme", 12], ["Globex", None]]},
            {"headers": ["x"], "rows": [[1.5]]},
        ], "leads")
        assert result["success"] is True
        assert result["sheets"] == [{"name": "Leads_ UAE", "rows": 2}, {"name": "Sheet2", "rows": 1}]
        wb = load_workbook(workspace / result["file_name"])
        ws = wb["Leads_ UAE"]
        assert ws["A1"].value == "company"
        assert ws["B2"].value == 12
        assert wb["Sheet2"]["A2"].value == 1.5

    async def test_saved_as_artifact(self, workspace):
        saved = {}

        class Ctx:
            def save_artifact(self, name, part):
                saved[name] = part.inline_data.mime_type

        result = await documents.write_spreadsheet(
            [{"headers": ["a"], "rows": [[1]]}], "t", tool_context=Ctx())
        assert saved[result["file_name"]].endswith("spreadsheetml.sheet")

    async def test_no_sheets(self, workspace):
        assert (await documents.write_spreadsheet([], "t"))["success"] is False


class TestTextFile:
    @pytest.mark.parametrize("kind", ["md", "txt", "csv", "html"])
    async def test_types(self, workspace, kind):
        result = await documents.write_text_file("post", "# Hello\nمرحبا", kind)
        assert result["success"] is True
        assert result["file_name"].endswith(f".{kind}")
        assert (workspace / result["file_name"]).read_text(encoding="utf-8").endswith("مرحبا")

    async def test_unknown_type(self, workspace):
        assert (await documents.write_text_file("x", "y", "exe"))["success"] is False

    async def test_empty_content(self, workspace):
        assert (await documents.write_text_file("x", "  "))["success"] is False
